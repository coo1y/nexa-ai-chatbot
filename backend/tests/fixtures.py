"""Builders for realistic test files (PDF, DOCX, XLSX, images)."""

import io

from PIL import Image


def make_pdf(pages: list[str]) -> bytes:
    """Minimal but valid multi-page PDF with extractable text."""
    objects: list[bytes] = []
    page_ids = []
    font_id = 3 + 2 * len(pages)
    for index, text in enumerate(pages):
        page_id, content_id = 3 + 2 * index, 4 + 2 * index
        page_ids.append(page_id)
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode()
        objects.append(
            f"{page_id} 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >> endobj\n".encode()
        )
        objects.append(
            f"{content_id} 0 obj << /Length {len(stream)} >> stream\n".encode() + stream + b"\nendstream endobj\n"
        )
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    header = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        f"2 0 obj << /Type /Pages /Kids [{kids}] /Count {len(pages)} >> endobj\n".encode(),
    ]
    font = f"{font_id} 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n".encode()
    body = header + objects + [font]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for obj in body:
        offsets.append(out.tell())
        out.write(obj)
    xref = out.tell()
    out.write(f"xref\n0 {len(body) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer << /Size {len(body) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx(paragraphs: list[str]) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "cell A"
    table.rows[0].cells[1].text = "cell B"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_xlsx(rows: list[list[object]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def make_image(
    width: int = 64, height: int = 48, fmt: str = "PNG", color: tuple = (200, 30, 30), exif: bool = False
) -> bytes:
    image = Image.new("RGB", (width, height), color)
    out = io.BytesIO()
    if exif:
        data = Image.Exif()
        data[0x010F] = "SecretCameraMaker"  # Make
        image.save(out, format=fmt, exif=data)
    else:
        image.save(out, format=fmt)
    return out.getvalue()
