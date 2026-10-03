"""Upload validation and content extraction.

Documents are reduced to plain text; images are validated, stripped of metadata (EXIF/GPS)
and downscaled to the size vision models actually use, which keeps requests fast.
"""

import asyncio
import io
import logging
import zipfile
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import PurePath
from typing import ClassVar

from PIL import Image, UnidentifiedImageError

from app.core.config import Settings
from app.core.errors import PayloadTooLarge, UnsupportedFile, ValidationFailed

logger = logging.getLogger(__name__)

Image.MAX_IMAGE_PIXELS = 50_000_000  # guard against decompression bombs

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".tsv",
    ".json",
    ".xml",
    ".yaml",
    ".yml",
    ".log",
    ".ini",
    ".toml",
    ".py",
    ".js",
    ".mjs",
    ".ts",
    ".tsx",
    ".jsx",
    ".java",
    ".kt",
    ".swift",
    ".c",
    ".h",
    ".cpp",
    ".hpp",
    ".cs",
    ".go",
    ".rs",
    ".rb",
    ".php",
    ".sql",
    ".sh",
    ".css",
    ".scss",
    ".r",
    ".m",
    ".scala",
    ".dart",
    ".lua",
}
HTML_EXTENSIONS = {".html", ".htm"}
OFFICE_EXTENSIONS = {".pdf", ".docx", ".xlsx"}
DOCUMENT_EXTENSIONS = sorted(TEXT_EXTENSIONS | HTML_EXTENSIONS | OFFICE_EXTENSIONS)
IMAGE_EXTENSIONS = [".png", ".jpg", ".jpeg", ".webp", ".gif"]

MAX_PDF_PAGES = 500
MAX_XLSX_ROWS_PER_SHEET = 5000
MAX_ZIP_UNCOMPRESSED = 100 * 1024 * 1024


@dataclass(slots=True)
class ProcessedFile:
    kind: str  # "document" | "image"
    content_type: str
    text: str | None = None
    page_count: int | None = None
    truncated: bool = False
    image_bytes: bytes | None = None
    width: int | None = None
    height: int | None = None


class FileService:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    @staticmethod
    def sanitize_filename(name: str | None) -> str:
        base = PurePath((name or "upload").replace("\\", "/")).name
        cleaned = "".join(ch for ch in base if ch.isprintable() and ch not in '<>:"|?*').strip(" .")
        return (cleaned or "upload")[:255]

    async def process(self, filename: str, data: bytes) -> ProcessedFile:
        if not data:
            raise ValidationFailed("The file is empty.")
        if len(data) > self.settings.upload_max_bytes:
            limit_mb = self.settings.upload_max_bytes // (1024 * 1024)
            raise PayloadTooLarge(f"The file is larger than the {limit_mb} MB limit.")
        ext = PurePath(filename).suffix.lower()
        if ext in IMAGE_EXTENSIONS:
            return await asyncio.to_thread(self._process_image, data)
        if ext in DOCUMENT_EXTENSIONS:
            processed = await asyncio.to_thread(self._process_document, ext, data)
            if not (processed.text or "").strip():
                raise ValidationFailed(
                    "No readable text was found in this document (scanned or image-only files are not supported)."
                )
            return processed
        raise UnsupportedFile(
            f"Unsupported file type '{ext or 'unknown'}'. Supported: documents "
            "(PDF, DOCX, XLSX, text, Markdown, CSV, JSON, HTML, code) and images (PNG, JPEG, WebP, GIF)."
        )

    # --- images -------------------------------------------------------------------

    def _process_image(self, data: bytes) -> ProcessedFile:
        try:
            with Image.open(io.BytesIO(data)) as probe:
                probe.verify()
            image: Image.Image = Image.open(io.BytesIO(data))
            image.seek(0)  # first frame of animated GIFs
            image.load()
        except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
            raise UnsupportedFile("The image could not be read. Please upload a valid PNG, JPEG, WebP or GIF.") from exc

        has_alpha = image.mode in ("RGBA", "LA") or (image.mode == "P" and "transparency" in image.info)
        image = image.convert("RGBA" if has_alpha else "RGB")
        image.thumbnail((self.settings.image_max_dimension, self.settings.image_max_dimension))
        out = io.BytesIO()
        # Re-encoding drops EXIF metadata (including GPS location).
        if has_alpha:
            image.save(out, format="PNG", optimize=True)
            content_type = "image/png"
        else:
            image.save(out, format="JPEG", quality=85, optimize=True)
            content_type = "image/jpeg"
        return ProcessedFile(
            kind="image", content_type=content_type, image_bytes=out.getvalue(), width=image.width, height=image.height
        )

    # --- documents ------------------------------------------------------------------

    def _process_document(self, ext: str, data: bytes) -> ProcessedFile:
        page_count: int | None = None
        if ext == ".pdf":
            text, page_count, content_type = *self._pdf(data), "application/pdf"
        elif ext == ".docx":
            text = self._docx(data)
            content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        elif ext == ".xlsx":
            text = self._xlsx(data)
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        elif ext in HTML_EXTENSIONS:
            text, content_type = _html_to_text(_decode_text(data)), "text/html"
        else:
            text = _decode_text(data)
            content_type = {".csv": "text/csv", ".json": "application/json", ".md": "text/markdown"}.get(
                ext, "text/plain"
            )
        text = text.replace("\x00", "")
        truncated = len(text) > self.settings.document_max_chars
        if truncated:
            text = text[: self.settings.document_max_chars]
        return ProcessedFile(
            kind="document", content_type=content_type, text=text, page_count=page_count, truncated=truncated
        )

    @staticmethod
    def _pdf(data: bytes) -> tuple[str, int]:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        if not data.startswith(b"%PDF"):
            raise UnsupportedFile("This file is not a valid PDF.")
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted and not reader.decrypt(""):
                raise UnsupportedFile("Password-protected PDFs are not supported.")
            pages = reader.pages
            parts = []
            for index, page in enumerate(pages[:MAX_PDF_PAGES], start=1):
                page_text = (page.extract_text() or "").strip()
                if page_text:
                    parts.append(f"--- Page {index} ---\n{page_text}")
            return "\n\n".join(parts), len(pages)
        except UnsupportedFile:
            raise
        except (PdfReadError, ValueError, KeyError, OSError) as exc:
            raise UnsupportedFile("The PDF could not be read.") from exc

    @staticmethod
    def _check_zip(data: bytes) -> None:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(info.file_size for info in archive.infolist()) > MAX_ZIP_UNCOMPRESSED:
                    raise UnsupportedFile("The document expands to an unsafe size.")
        except zipfile.BadZipFile as exc:
            raise UnsupportedFile("The document is corrupted or not a valid Office file.") from exc

    def _docx(self, data: bytes) -> str:
        import docx

        self._check_zip(data)
        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:
            raise UnsupportedFile("The Word document could not be read.") from exc
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)

    def _xlsx(self, data: bytes) -> str:
        from openpyxl import load_workbook

        self._check_zip(data)
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise UnsupportedFile("The spreadsheet could not be read.") from exc
        parts = []
        for sheet in workbook.worksheets:
            parts.append(f"--- Sheet: {sheet.title} ---")
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_XLSX_ROWS_PER_SHEET:
                    parts.append("... (rows truncated)")
                    break
                if any(cell is not None for cell in row):
                    parts.append(",".join("" if cell is None else str(cell) for cell in row))
        workbook.close()
        return "\n".join(parts)


def _decode_text(data: bytes) -> str:
    if data.count(b"\x00") > len(data) // 100:
        raise UnsupportedFile("This looks like a binary file, not a text document.")
    for encoding in ("utf-8-sig", "utf-16" if data[:2] in (b"\xff\xfe", b"\xfe\xff") else "utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


class _TextExtractor(HTMLParser):
    _SKIP: ClassVar[set[str]] = {"script", "style", "noscript", "template", "svg"}
    _BLOCK: ClassVar[set[str]] = {
        "p",
        "div",
        "br",
        "li",
        "tr",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "section",
        "article",
        "pre",
    }

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list) -> None:  # type: ignore[type-arg]
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP and self._skip_depth:
            self._skip_depth -= 1

    def handle_data(self, data: str) -> None:
        if not self._skip_depth:
            self.parts.append(data)


def _html_to_text(markup: str) -> str:
    parser = _TextExtractor()
    parser.feed(markup)
    lines = [" ".join(line.split()) for line in "".join(parser.parts).splitlines()]
    return "\n".join(line for line in lines if line)
